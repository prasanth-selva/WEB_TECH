from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
DOCX = ROOT / 'Assignment-Documentation.docx'
NAVY = '142137'
MINT = '57E6B1'
TEAL = '168B67'
PALE = 'EAFBF4'
PALE_BLUE = 'F1F5F8'
MID = '526176'
LINE = 'DFE6EB'
WHITE = 'FFFFFF'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tc_pr.append(shd)


def cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr(); mar = OxmlElement('w:tcMar')
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        tag = OxmlElement(f'w:{side}'); tag.set(qn('w:w'), str(value)); tag.set(qn('w:type'), 'dxa'); mar.append(tag)
    tc_pr.append(mar)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr(); tbl_header = OxmlElement('w:tblHeader'); tbl_header.set(qn('w:val'), 'true'); tr_pr.append(tbl_header)


def set_cell_borders(table, color=LINE, size='5'):
    tbl_pr = table._tbl.tblPr; borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = OxmlElement(f'w:{edge}'); tag.set(qn('w:val'), 'single'); tag.set(qn('w:sz'), size); tag.set(qn('w:color'), color); borders.append(tag)
    tbl_pr.append(borders)


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('PULSEBOARD  ·  '); run.font.name = 'Aptos'; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string('8995A3')
    run = paragraph.add_run('PAGE '); run.font.name = 'Aptos'; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string('8995A3')
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run = paragraph.add_run(); run._r.extend([begin, instr, separate, text, end])
    run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string('8995A3')


def set_cell_text(cell, text, bold=False, color=NAVY, size=9, font='Aptos'):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(text)); r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)


def table(doc, headers, rows, widths=None, font_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=WHITE, size=font_size)
        shade(t.rows[0].cells[i], NAVY)
    set_repeat_table_header(t.rows[0])
    for row_num, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, color=MID, size=font_size)
            if row_num % 2 == 1: shade(cells[i], 'F7F9FA')
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths): row.cells[i].width = Inches(width)
    set_cell_borders(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def para(doc, text='', style=None, size=10.3, color=NAVY, bold=False, italic=False, after=6, before=0, align=None, font='Aptos', line=1.15):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before); p.paragraph_format.line_spacing = line
    if align is not None: p.alignment = align
    r = p.add_run(text); r.font.name = font; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color); r.bold = bold; r.italic = italic
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5 if level == 1 else 3)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(.22); p.paragraph_format.first_line_indent = Inches(-.12)
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text); r.font.name = 'Aptos'; r.font.size = Pt(9.7); r.font.color.rgb = RGBColor.from_string(MID)
    return p


def code_block(doc, text, font_size=8.1):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    cell = t.cell(0,0); shade(cell, 'F3F6F8'); cell_margins(cell, top=120, start=150, bottom=120, end=150)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text); r.font.name = 'Consolas'; r.font.size = Pt(font_size); r.font.color.rgb = RGBColor.from_string('314157')
    set_cell_borders(t, color='E0E6EA', size='4')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def page_break(doc): doc.add_page_break()


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
sec.top_margin = Inches(.68); sec.bottom_margin = Inches(.68); sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
sec.header_distance = Inches(.32); sec.footer_distance = Inches(.32)
styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Aptos'; normal.font.size = Pt(10.3); normal.font.color.rgb = RGBColor.from_string(NAVY)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.15
for name, size, color in [('Title',32,NAVY),('Heading 1',19,NAVY),('Heading 2',13,TEAL),('Heading 3',11,NAVY)]:
    st = styles[name]; st.font.name = 'Aptos Display' if name in ('Title','Heading 1') else 'Aptos'; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.keep_with_next = True
header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.LEFT
rr = header.add_run('WEB TECHNOLOGY  /  FULL-STACK APPLICATION'); rr.font.name = 'Aptos'; rr.font.size = Pt(8); rr.font.bold = True; rr.font.color.rgb = RGBColor.from_string('8995A3')
set_page_number(sec.footer.paragraphs[0])

# Cover page
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(72); p.paragraph_format.space_after = Pt(14)
r = p.add_run('REAL-TIME WEB APPLICATION  ·  ASSIGNMENT REPORT'); r.font.name='Aptos'; r.font.size=Pt(10); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL)
p = doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.0
r=p.add_run('PulseBoard'); r.font.name='Aptos Display'; r.font.size=Pt(42); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
para(doc, 'A live classroom task board powered by REST APIs, WebSockets and SQLite', size=17, color=MID, after=19, line=1.2)
# Accent rule
rule = doc.add_paragraph(); rule.paragraph_format.space_after=Pt(18)
pr=rule._p.get_or_add_pPr(); p_bdr=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'22'); bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),MINT); p_bdr.append(bottom); pr.append(p_bdr)
para(doc, 'PROJECT SUMMARY', size=9, color=TEAL, bold=True, after=5)
para(doc, 'PulseBoard lets classmates create, review and update shared tasks. Persistent changes travel through a JSON REST API into SQLite; authenticated Socket.IO connections distribute each change to the whole classroom without a refresh.', size=12, color=NAVY, after=20, line=1.35)
feature = doc.add_table(rows=1, cols=3); feature.alignment=WD_TABLE_ALIGNMENT.CENTER; feature.autofit=False
features=[('01  REST + JSON','GET · POST · PUT · DELETE'),('02  LIVE EVENTS','Socket.IO · shared class room'),('03  DATA + AUTH','SQLite · bcrypt · JWT')]
for i,(title,sub) in enumerate(features):
    c=feature.cell(0,i); c.width=Inches(2.2); shade(c, NAVY); cell_margins(c,top=190,start=155,bottom=190,end=155)
    c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(4); rr=p.add_run(title); rr.font.name='Aptos'; rr.font.size=Pt(9); rr.font.bold=True; rr.font.color.rgb=RGBColor.from_string(MINT)
    p2=c.add_paragraph(); p2.paragraph_format.space_after=Pt(0); rr=p2.add_run(sub); rr.font.name='Aptos'; rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string(WHITE)
set_cell_borders(feature,color=NAVY,size='0')
para(doc, 'SUBMISSION DETAILS', size=9, color=TEAL, bold=True, before=24, after=7)
table(doc,['Item','Details'],[
    ('Assignment','Real-time web application using REST APIs and WebSockets'),
    ('Application','PulseBoard — collaborative classroom task manager'),
    ('Repository','github.com/prasanth-selva/WEB_TECH'),
    ('Prepared','07 October 2026'),
],[1.35,5.2],font_size=9)
para(doc, 'Demonstration account  ·  demo@pulseboard.app  /  Classroom2026!', size=9, color='7F8B99', before=6, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)

# Contents and purpose
page_break(doc)
heading(doc,'Contents',1)
contents=[
 ('1','Project overview and requirements'),('2','System architecture and workflow'),('3','REST API design and JSON examples'),('4','Database design and real-time exchange'),('5','Authentication, security and error handling'),('6','Verification results and rubric mapping'),('7','Application screenshots'),('8','Run instructions, project structure and conclusion')]
table(doc,['Section','Topic'],contents,[.85,5.7],font_size=9)
heading(doc,'1. Project overview',1)
para(doc,'PulseBoard is a browser-based collaboration tool for a single demonstration classroom. Each student signs in, sees the same shared task board, and can create, update, move or remove tasks. The project deliberately combines a request/response interface for durable data with a push channel for immediate updates.',size=10.2)
heading(doc,'Assignment objectives',2)
for text in [
 'Demonstrate client–server communication using REST endpoints and JSON request/response bodies.',
 'Implement the complete task lifecycle with GET, POST, PUT and DELETE operations backed by SQLite.',
 'Authenticate HTTP and WebSocket clients, validate input, and return predictable errors.',
 'Broadcast task changes to separately authenticated classmates in real time.'
]: bullet(doc,text)
heading(doc,'Scope decision',2)
para(doc,'The demo represents one shared class, “Class 04 · Web Tech”. Any valid account registered in this teaching instance joins that room and can collaborate on its tasks. The task creator remains visible for attribution. A production product should add explicit classroom membership, role-based permissions and protected multi-tenant room selection.',size=9.8,color=MID)

# Architecture
page_break(doc)
heading(doc,'2. System architecture and workflow',1)
para(doc,'The browser uses two complementary channels. Fetch requests carry a bearer JWT to Express and receive JSON responses; a Socket.IO connection authenticated with the same token subscribes to the classroom room. Express writes the task change to SQLite before emitting the corresponding event.',size=10)
fig=doc.add_paragraph(); fig.alignment=WD_ALIGN_PARAGRAPH.CENTER; fig.paragraph_format.space_after=Pt(2)
run=fig.add_run(); picture=run.add_picture(str(ROOT/'architecture.png'),width=Inches(6.45)); picture._inline.docPr.set('descr','PulseBoard architecture: browser, Express REST API, SQLite, authenticated Socket.IO and shared classroom events.')
para(doc,'Figure 1. PulseBoard client–server architecture.',size=8.5,color='788594',italic=True,after=9,align=WD_ALIGN_PARAGRAPH.CENTER)
table(doc,['Component','Responsibility'],[
 ('Browser UI','Semantic HTML/CSS; fetch-based CRUD; token-bearing Socket.IO connection; board and live activity feed.'),
 ('Express server','Hosts static assets, authenticates requests, validates fields, returns JSON and centralizes errors.'),
 ('SQLite','Persists account and task records; foreign keys, constraints and indexes protect/accelerate data access.'),
 ('Socket.IO','Authenticates the handshake, joins the shared class room, broadcasts task and presence events.'),
],[1.45,5.1],font_size=8.6)
heading(doc,'End-to-end workflow',2)
for text in [
 'Sign in or register: bcrypt verifies or hashes the password; the server returns an eight-hour JWT.',
 'Load the board: the client sends GET /api/tasks with Authorization: Bearer <token>; Express reads SQLite and returns JSON.',
 'Mutate a task: POST, PUT or DELETE persists the change and returns the resulting representation/status.',
 'Broadcast: after the database write, Socket.IO emits task:created, task:updated or task:deleted to all authenticated room members.',
 'Synchronize: each connected browser updates its task cards, counts and activity feed without reloading.'
]: bullet(doc,text)

# REST API
page_break(doc)
heading(doc,'3. REST API design and JSON exchange',1)
para(doc,'All task routes except the health check require a valid bearer token. JSON responses use one success envelope and one error envelope. IDs are positive integers; task text, status, priority and date values are validated before SQL is executed.',size=10)
table(doc,['Method','Endpoint','Purpose','Success'],[
 ('POST','/api/auth/register','Validate account fields, hash password, create account and issue JWT.','201'),
 ('POST','/api/auth/login','Verify credentials and issue JWT.','200'),
 ('GET','/api/auth/me','Resolve the token’s account identity.','200'),
 ('GET','/api/tasks','List shared tasks; optional status and q filters.','200'),
 ('POST','/api/tasks','Create a task and broadcast task:created.','201'),
 ('PUT','/api/tasks/:id','Update task fields/status and broadcast task:updated.','200'),
 ('DELETE','/api/tasks/:id','Delete a task and broadcast task:deleted.','200'),
 ('GET','/api/health','Check server and SQLite connectivity.','200'),
],[.58,1.55,3.75,.62],font_size=7.8)
heading(doc,'Example: create a task',2)
code_block(doc,'POST /api/tasks\nAuthorization: Bearer <JWT>\nContent-Type: application/json\n\n{\n  "title": "Prepare the lab worksheet",\n  "description": "Add diagrams and upload the answer key.",\n  "status": "todo",\n  "priority": "high",\n  "dueDate": null\n}',font_size=8.1)
para(doc,'A successful response has HTTP 201 and includes the saved row, its generated ID, timestamps and creator attribution:',size=9.4,after=4)
code_block(doc,'{ "success": true, "data": { "task": {\n  "id": 24, "title": "Prepare the lab worksheet",\n  "status": "todo", "priority": "high",\n  "createdBy": "Alex Morgan", "creatorId": 1\n} } }',font_size=8.1)
para(doc,'A successful task update returns HTTP 200; deletion returns `{ "deleted": true, "id": 24 }`. Invalid fields return HTTP 400, missing resources return 404, unauthenticated requests return 401, duplicate emails return 409, and unexpected failures return 500.',size=9.3,color=MID)

# Database and realtime
page_break(doc)
heading(doc,'4. Database design and real-time updates',1)
para(doc,'SQLite is the persistent system of record. Prepared statements are used for reads and writes; task rows retain the account ID of their creator for attribution. The shared board returns all class tasks to authenticated accounts.',size=10)
table(doc,['Table','Column','Type / rule','Purpose'],[
 ('users','id','INTEGER PRIMARY KEY','Unique account ID.'),
 ('users','name, email','TEXT; name 2–60; email UNIQUE','Profile and sign-in identity.'),
 ('users','password_hash','TEXT NOT NULL','bcrypt digest; plaintext is never stored.'),
 ('tasks','id, user_id','INTEGER; user_id FK → users','Task ID and creator account.'),
 ('tasks','title, description','TEXT; 3–120 / max 500','Task content with database constraints.'),
 ('tasks','status, priority','CHECK-constrained TEXT','todo/doing/done; low/normal/high.'),
 ('tasks','due_date, timestamps','TEXT / ISO-like date; UTC timestamps','Optional due date and audit timestamps.'),
],[.72,1.35,2.3,2.15],font_size=7.7)
heading(doc,'Socket.IO event contract',2)
para(doc,'The browser provides its JWT in the Socket.IO handshake. Invalid or expired tokens are rejected before the client joins `class:04-web-tech`. Server events are emitted after the REST mutation has been persisted.',size=9.7)
table(doc,['Event','Payload fields','Client effect'],[
 ('task:created','task, actor, at','Insert the new card and activity entry.'),
 ('task:updated','task, actor, at','Replace the affected card and refresh counts.'),
 ('task:deleted','id, title, actor, at','Remove the card and log the event.'),
 ('presence:count','count','Update active classroom connection count.'),
],[1.25,2.55,2.72],font_size=8.3)
code_block(doc,'{ "task": { "id": 24, "title": "Prepare the lab worksheet",\n  "status": "todo", "priority": "high" },\n  "actor": "Alex Morgan", "at": "2026-10-07T09:32:20.965Z" }',font_size=8.0)
para(doc,'The activity feed in the browser reflects live events in memory; it is not a separate persistent audit-log table in this demo.',size=9,color='7B8796',italic=True)

# Security and errors
page_break(doc)
heading(doc,'5. Authentication, security and error handling',1)
table(doc,['Control','Implementation in PulseBoard'],[
 ('Password handling','bcryptjs hash with work factor 10 at registration; compare hash during login.'),
 ('Session','JWT signed by JWT_SECRET with an eight-hour expiry; bearer header on REST and handshake token on Socket.IO.'),
 ('Request hardening','Helmet security headers; body size limited to 32 KB; authentication routes rate-limited.'),
 ('Validation','Email syntax, name/password bounds, title/description lengths, allowed status/priority, date format and positive IDs.'),
 ('Database safety','Parameterized SQL, foreign keys, schema CHECK constraints, and token subject resolution to a registered user.'),
 ('Browser output','User-entered titles/descriptions are assigned with textContent; HTML injection is avoided.'),
 ('Failure behavior','Central JSON error envelope, meaningful 400/401/404/409/500 status codes and readable UI toast/form errors.'),
],[1.45,5.1],font_size=8.4)
heading(doc,'Error response examples',2)
code_block(doc,'{ "success": false, "error": {\n  "code": "VALIDATION_ERROR",\n  "message": "Title must be 3–120 characters."\n} }',font_size=8.2)
heading(doc,'Operational note',2)
para(doc,'Set a long, random JWT_SECRET before deployment. The built-in fallback is for local teaching use only. HTTPS, secure secret storage, membership-aware class rooms, stronger account recovery, audit persistence, backups, and Socket.IO scaling are intentionally outside this single-process assignment demo.',size=9.8)

# Verification and rubric
page_break(doc)
heading(doc,'6. Verification results and rubric mapping',1)
para(doc,'The Node.js integration suite starts an isolated HTTP server and exercises the application over HTTP and a genuine Socket.IO websocket. The final run passed all four test groups.',size=10)
table(doc,['Test group','Coverage','Result'],[
 ('Health + storage','HTTP health response and SQLite probe.','PASS'),
 ('Auth','Register/login; bearer-token protection; wrong-password rejection.','PASS'),
 ('REST CRUD + validation','GET/POST/PUT/DELETE, filters, malformed JSON, field validation and 404 behavior.','PASS'),
 ('Shared collaboration','Cross-account board access/edit and a task event delivered to another authenticated account.','PASS'),
],[1.35,4.35,.85],font_size=8.2)
para(doc,'Command: `npm test`  ·  Final run: 4 passed, 0 failed.',size=9.2,color=TEAL,bold=True,after=11)
heading(doc,'Assessment rubric alignment',2)
table(doc,['Criterion','Maximum','Project evidence'],[
 ('System design & architecture','8','Architecture figure, request/event workflow, component responsibilities and documented API routes.'),
 ('RESTful web service','14','Authenticated GET/POST/PUT/DELETE task endpoints, filters, HTTP status codes and JSON error contract.'),
 ('Database & JSON exchange','10','SQLite users/tasks schema, constraints, persistence, creator attribution and JSON examples.'),
 ('Advanced technologies, security & errors','8','JWT + bcrypt, Socket.IO room broadcasts, validation, Helmet, rate limiting and central error handling.'),
 ('Total','40','This table maps evidence to criteria; it does not award a grade.'),
],[2.55,.72,3.28],font_size=7.9)
heading(doc,'Known limitations',2)
para(doc,'The app models one shared classroom and one server process. Users who can register join the same room; authorization is therefore suitable for a controlled classroom demonstration, not a public multi-class service. The report screenshots show a real task created from a second signed-in account and delivered to the first account.',size=9.4,color=MID)

# Screenshots
page_break(doc)
heading(doc,'7. Application screenshots',1)
para(doc,'The following figures were captured from the running web application. The live board image shows the shared classroom UI after a second authenticated account creates a task; the first browser displays the activity update and task card without reloading.',size=9.8)
for image, caption, alt in [
 ('screenshots/01-sign-in.png','Figure 2. Sign-in screen with account creation and the local demo access option.','PulseBoard sign-in screen.'),
 ('screenshots/02-live-classroom-board.png','Figure 3. Shared classroom board with REST-backed task lanes and a Socket.IO activity update from another account.','PulseBoard live classroom dashboard showing task lanes and an activity event.')
]:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
    r=p.add_run(); picture=r.add_picture(str(ROOT/image),width=Inches(5.75)); picture._inline.docPr.set('descr',alt)
    para(doc,caption,size=8.3,color='788594',italic=True,after=7,align=WD_ALIGN_PARAGRAPH.CENTER)

# Run and conclusion
page_break(doc)
heading(doc,'8. Run instructions, project structure and conclusion',1)
heading(doc,'Run locally',2)
code_block(doc,'Requirements: Node.js 20+ and npm\n\nnpm install\nexport JWT_SECRET="replace-with-a-long-random-secret"\nnpm start\n\nOpen http://localhost:3000\nDemo: demo@pulseboard.app / Classroom2026!',font_size=8.3)
para(doc,'The database is created automatically at `data/pulseboard.sqlite`; first startup creates the demo account and three sample tasks. Run `npm test` to execute the integration suite.',size=9.6)
heading(doc,'Project structure',2)
table(doc,['Path','Role'],[
 ('server.js','REST endpoints, JWT middleware, Socket.IO room, static hosting and error handling.'),
 ('src/database.js','SQLite schema, initialization, seed data and database connection.'),
 ('public/index.html · styles.css · app.js','UI, responsive styles, fetch client, CRUD and socket synchronization.'),
 ('test/api.test.js','HTTP/WebSocket integration tests.'),
 ('docs/','Assignment report, screenshots, diagram and report builder.'),
 ('README.md · plan.md','Setup, architecture notes and implementation plan.'),
],[2.75,3.8],font_size=8.2)
heading(doc,'Conclusion',2)
para(doc,'PulseBoard demonstrates the full request-and-push cycle expected from a real-time web application: authenticated JSON requests persist shared classroom tasks in SQLite, and authorized WebSocket clients receive the committed change immediately. The implementation covers architecture, CRUD, data handling, authentication, validation and error feedback, with automated integration evidence and screenshots included for assessment.',size=10.1,line=1.25)
para(doc,'Deliverables: runnable source in the WEB_TECH repository, this DOCX report, and its PDF export.',size=9,color=TEAL,bold=True,before=4)

doc.core_properties.title = 'PulseBoard — Real-Time Web Application Assignment Documentation'
doc.core_properties.subject = 'REST APIs, WebSockets, authentication, SQLite and JSON data exchange'
doc.core_properties.author = 'PulseBoard Project'
doc.core_properties.keywords = 'REST, WebSocket, Socket.IO, SQLite, JWT, assignment'
doc.save(DOCX)
print(DOCX)
